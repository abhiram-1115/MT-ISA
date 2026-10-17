"""
Analyse the label audit of the implicit test rows.

Inputs : filled annotation sheets (--annotators a.xlsx b.xlsx ...), audit/key.csv, LLM labels
         (audit/labels_qwen7b.csv), test_predictions.json of the runs in models/ and models_laptop/
Output : results/audit_report.txt

Usage: python audit/analyze.py --annotators audit/ann_A.xlsx audit/ann_B.xlsx audit/ann_C.xlsx
"""
import argparse
import csv
import glob
import json
import os
import re
import statistics as st
import sys
from collections import Counter, defaultdict
from itertools import combinations

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from guideline import DATASETS, LABELS, SHEET  # noqa: E402

DS_MODELS = {"restaurant": "models", "laptop": "models_laptop"}
OUT = []


def log(s=""):
    print(s)
    OUT.append(s)


def pct(a, b):
    return f"{100 * a / b:5.1f}% ({a}/{b})" if b else "   n/a (0)"


# ----------------------------------------------------------------------------
# agreement statistics
# ----------------------------------------------------------------------------
def cohen_kappa(a, b):
    n = len(a)
    if n == 0:
        return float("nan")
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[l] * cb[l] for l in LABELS) / (n * n)
    return (po - pe) / (1 - pe) if pe < 1 else float("nan")


def fleiss_kappa(table):
    """table: list of per-item label lists, every item rated by the same number of raters."""
    if not table:
        return float("nan")
    m = len(table[0])
    counts = np.array([[sum(1 for x in row if x == l) for l in LABELS] for row in table], float)
    N = len(table)
    p_j = counts.sum(0) / (N * m)
    P_i = ((counts ** 2).sum(1) - m) / (m * (m - 1))
    P_bar, P_e = P_i.mean(), (p_j ** 2).sum()
    return float((P_bar - P_e) / (1 - P_e)) if P_e < 1 else float("nan")


def majority(labels):
    """Plurality label; a tie for the top count -> None ('no majority')."""
    c = Counter(l for l in labels if l)
    if not c:
        return None
    top = c.most_common()
    if len(top) > 1 and top[0][1] == top[1][1]:
        return None
    return top[0][0]


def macro_f1(g, p):
    g, p = np.asarray(g), np.asarray(p)
    f = []
    for c in range(3):
        tp = np.sum((g == c) & (p == c)); fp = np.sum((g != c) & (p == c)); fn = np.sum((g == c) & (p != c))
        d = 2 * tp + fp + fn
        f.append(2 * tp / d if d else 0.0)
    return float(np.mean(f))


# ----------------------------------------------------------------------------
# loading
# ----------------------------------------------------------------------------
def read_sheet(path):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.worksheets[-1]
    it = ws.iter_rows(values_only=True)
    header = [str(h).strip().lower() if h else "" for h in next(it)]
    ci = {h: header.index(h) for h in ("item", "label", "sure", "note") if h in header}
    out, bad = {}, 0
    for r in it:
        if not r or r[ci["item"]] is None:
            continue
        item = str(r[ci["item"]]).strip()
        lab = str(r[ci["label"]] or "").strip().lower()
        if lab and lab not in LABELS:
            bad += 1
            lab = ""
        sure = str(r[ci["sure"]] or "").strip().lower() if "sure" in ci else ""
        out[item] = {"label": lab, "sure": sure,
                     "note": str(r[ci["note"]] or "") if "note" in ci else ""}
    return out, bad


def load_runs(models_dir, lrs):
    """name -> lr -> seed -> {instance_id: pred}"""
    runs = defaultdict(lambda: defaultdict(dict))
    for f in glob.glob(os.path.join(models_dir, "exp_*", "test_predictions.json")):
        m = re.match(r"exp_(.+)_lr([^_]+)_s(\d+)$", os.path.basename(os.path.dirname(f)))
        if not m or m.group(2) not in lrs:
            continue
        rows = json.load(open(f, encoding="utf-8"))
        runs[m.group(1)][m.group(2)][int(m.group(3))] = {r["id"]: r["pred"] for r in rows}
    return runs


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--annotators", nargs="+", required=True, help="filled annotation sheets (.xlsx)")
    ap.add_argument("--key", default="audit/key.csv")
    ap.add_argument("--llm", default="audit/labels_qwen7b.csv")
    ap.add_argument("--llm-name", default="Qwen2.5-7B")
    ap.add_argument("--lrs", default="1e-4,2e-4", help="lr tags of runs used for model difficulty")
    ap.add_argument("--rescore", default="polonly_all_gv:1e-4,ce_all_gv:1e-4,mtisa_in_all_gv:1e-4,lora_pol:2e-4")
    ap.add_argument("--n-list", type=int, default=30)
    ap.add_argument("--out", default="results/audit_report.txt")
    args = ap.parse_args()

    # ---- key + texts ----
    key = {}
    with open(args.key, encoding="utf-8") as f:
        for r in csv.DictReader(f):
            key[r["item"]] = {"dataset": r["dataset"], "id": r["instance_id"], "gold": r["gold"].lower()}
    text = {}
    for ds, (impl, _) in DATASETS.items():
        if os.path.exists(impl):
            for r in json.load(open(impl, encoding="utf-8")):
                text[(ds, r["id"])] = (r["sentence"], r["target"])
    items = list(key)
    by_ds = {ds: [i for i in items if key[i]["dataset"] == ds] for ds in DATASETS}

    # ---- annotators ----
    ann = {}
    for p in args.annotators:
        name = os.path.splitext(os.path.basename(p))[0]
        if name in ann:
            name = p
        ann[name], bad = read_sheet(p)
        unknown = [i for i in ann[name] if i not in key]
        if bad or unknown:
            print(f"WARNING {name}: {bad} invalid labels ignored, {len(unknown)} unknown items")
    names = list(ann)
    lab = {n: {i: ann[n].get(i, {}).get("label", "") for i in items} for n in names}

    llm = {}
    if args.llm and os.path.exists(args.llm):
        with open(args.llm, encoding="utf-8") as f:
            llm = {r["item"]: (r.get("label") or "").strip().lower() for r in csv.DictReader(f)}
        llm = {i: l for i, l in llm.items() if l in LABELS}

    maj = {i: majority([lab[n][i] for n in names]) for i in items}
    gold = {i: key[i]["gold"] for i in items}

    log(f"LABEL AUDIT | {len(items)} rows ({', '.join(f'{ds} {len(v)}' for ds, v in by_ds.items())}) | "
        f"annotators: {', '.join(names)} | LLM: {args.llm_name} ({len(llm)} labels)")
    log()

    # ================= 1. per annotator =================
    log("=" * 90 + "\n1. PER ANNOTATOR: agreement with gold\n" + "=" * 90)
    log(f"{'annotator':<16}{'labelled':<10}{'overall':<20}" + "".join(f"{'gold ' + l:<20}" for l in LABELS)
        + "".join(f"{ds:<20}" for ds in DATASETS) + "sure=no")
    raters = [(n, lab[n]) for n in names] + ([(args.llm_name, llm)] if llm else [])
    for n, L in raters:
        done = [i for i in items if L.get(i)]
        row = f"{n:<16}{len(done):<10}"
        row += f"{pct(sum(L[i] == gold[i] for i in done), len(done)):<20}"
        for l in LABELS:
            s = [i for i in done if gold[i] == l]
            row += f"{pct(sum(L[i] == gold[i] for i in s), len(s)):<20}"
        for ds in DATASETS:
            s = [i for i in done if key[i]["dataset"] == ds]
            row += f"{pct(sum(L[i] == gold[i] for i in s), len(s)):<20}"
        if n in ann:
            sn = sum(1 for i in done if ann[n][i]["sure"] == "no")
            row += pct(sn, len(done))
        log(row)
    log("(per class = agreement on the rows with that gold label)")
    log()

    # ================= 2. inter-annotator agreement =================
    log("=" * 90 + "\n2. AGREEMENT AMONG HUMANS\n" + "=" * 90)
    if len(names) >= 2:
        full = [i for i in items if all(lab[n][i] for n in names)]
        log(f"Fleiss kappa ({len(names)} annotators, {len(full)} rows labelled by all): "
            f"{fleiss_kappa([[lab[n][i] for n in names] for i in full]):.3f}")
        for a, b in combinations(names, 2):
            both = [i for i in items if lab[a][i] and lab[b][i]]
            log(f"Cohen kappa {a} vs {b}: {cohen_kappa([lab[a][i] for i in both], [lab[b][i] for i in both]):.3f} "
                f"(n={len(both)}, raw agreement {pct(sum(lab[a][i] == lab[b][i] for i in both), len(both))})")
        if llm:
            for a in names:
                both = [i for i in items if lab[a][i] and llm.get(i)]
                log(f"Cohen kappa {a} vs {args.llm_name}: "
                    f"{cohen_kappa([lab[a][i] for i in both], [llm[i] for i in both]):.3f} (n={len(both)})")
    else:
        log("only one annotator: no inter-annotator agreement (the 'majority' is that annotator's label)")
    n_maj = sum(1 for i in items if maj[i])
    log(f"human majority label available: {pct(n_maj, len(items))}; no majority (tie or unlabelled): "
        f"{len(items) - n_maj}")
    log("majority label distribution: " + str(dict(Counter(maj[i] for i in items if maj[i]))))
    log()

    # ================= 3. majority / LLM vs gold =================
    log("=" * 90 + "\n3. LABEL DISAGREEMENT WITH GOLD\n" + "=" * 90)

    def disagree_block(title, L):
        log(title)
        sel = [i for i in items if L.get(i)]
        log(f"  overall            : {pct(sum(L[i] != gold[i] for i in sel), len(sel))}")
        for l in LABELS:
            s = [i for i in sel if gold[i] == l]
            moved = Counter(L[i] for i in s if L[i] != gold[i])
            log(f"  gold {l:<13} : {pct(sum(L[i] != gold[i] for i in s), len(s))}  relabelled as {dict(moved)}")
        for ds in DATASETS:
            s = [i for i in sel if key[i]["dataset"] == ds]
            log(f"  {ds:<18} : {pct(sum(L[i] != gold[i] for i in s), len(s))}")
    disagree_block("human majority != gold (rows with a majority):", {i: maj[i] for i in items if maj[i]})
    if llm:
        disagree_block(f"{args.llm_name} != gold:", llm)
    log()

    # ================= 4. by model difficulty =================
    log("=" * 90 + "\n4. BY MODEL DIFFICULTY (error rate over all runs, lr " + args.lrs + ")\n" + "=" * 90)
    lrs = set(args.lrs.split(","))
    runs_by_ds = {ds: load_runs(DS_MODELS[ds], lrs) for ds in DATASETS}
    err, nruns = {}, {}
    for i in items:
        ds, iid = key[i]["dataset"], key[i]["id"]
        preds = [d[iid] for name in runs_by_ds[ds].values() for lr in name.values()
                 for d in lr.values() if iid in d]
        if preds:
            err[i] = float(np.mean([p != gold[i] for p in preds]))
            nruns[i] = len(preds)
    for ds in DATASETS:
        n = sum(len(s) for name in runs_by_ds[ds].values() for s in name.values())
        log(f"{ds}: {n} runs in {DS_MODELS[ds]}/")
    bucket = {i: ("always-wrong" if e >= 0.9 else "always-right" if e <= 0.1 else "in-between")
              for i, e in err.items()}
    log(f"{'bucket':<14}{'n':<6}{'maj != gold':<20}{args.llm_name + ' != gold':<22}"
        f"{'mean annot. agree':<20}{'sure=no':<18}{'no majority'}")
    for b in ("always-wrong", "in-between", "always-right"):
        s = [i for i in items if bucket.get(i) == b]
        sm = [i for i in s if maj[i]]
        sl = [i for i in s if llm.get(i)]
        agree = [np.mean([lab[n][i] == gold[i] for n in names if lab[n][i]]) for i in s
                 if any(lab[n][i] for n in names)]
        sure_no = sum(1 for n in names for i in s if lab[n][i] and ann[n][i]["sure"] == "no")
        sure_all = sum(1 for n in names for i in s if lab[n][i])
        log(f"{b:<14}{len(s):<6}{pct(sum(maj[i] != gold[i] for i in sm), len(sm)):<20}"
            f"{pct(sum(llm[i] != gold[i] for i in sl), len(sl)):<22}"
            f"{(f'{100 * np.mean(agree):.1f}%' if agree else 'n/a'):<20}{pct(sure_no, sure_all):<18}"
            f"{len(s) - len(sm)}")
    log(f"(rows without any run: {len(items) - len(err)})")
    log()

    # ================= 5. re-scoring =================
    log("=" * 90 + "\n5. RE-SCORING MODELS (mean +/- std over seeds)\n" + "=" * 90)
    log("A = gold (all rows) | B = human majority (rows with a majority) | "
        "C = gold, only rows where the human majority agrees with gold")
    for ds in DATASETS:
        ds_items = by_ds[ds]
        sets = {"A": [i for i in ds_items],
                "B": [i for i in ds_items if maj[i]],
                "C": [i for i in ds_items if maj[i] and maj[i] == gold[i]]}
        ref = {"A": gold, "B": maj, "C": gold}
        log(f"--- {ds} (n: A {len(sets['A'])}, B {len(sets['B'])}, C {len(sets['C'])}) ---")
        log(f"{'config':<20}{'lr':<7}{'seeds':<6}" + "".join(f"{s + ' acc':<18}{s + ' F1':<18}" for s in "ABC"))
        for spec in args.rescore.split(","):
            cfg, lr = spec.split(":")
            seeds = runs_by_ds[ds].get(cfg, {}).get(lr, {})
            if not seeds:
                log(f"{cfg:<20}{lr:<7}{0:<6}(no runs)")
                continue
            row = f"{cfg:<20}{lr:<7}{len(seeds):<6}"
            for s in "ABC":
                accs, f1s = [], []
                for d in seeds.values():
                    sel = [i for i in sets[s] if key[i]["id"] in d]
                    if not sel:
                        continue
                    g = [LABELS.index(ref[s][i]) for i in sel]
                    p = [LABELS.index(d[key[i]["id"]]) for i in sel]
                    accs.append(float(np.mean(np.array(g) == np.array(p))))
                    f1s.append(macro_f1(g, p))
                fmt = (lambda v: f"{st.mean(v):.4f}+/-{st.stdev(v):.4f}" if len(v) > 1
                       else (f"{v[0]:.4f}" if v else "n/a"))
                row += f"{fmt(accs):<18}{fmt(f1s):<18}"
            log(row)
        log()

    # ================= 6. unanimous disagreement with gold =================
    log("=" * 90 + f"\n6. ROWS WHERE ALL HUMANS AGREE AND DISAGREE WITH GOLD (up to {args.n_list})\n" + "=" * 90)
    unan = [i for i in items if len(names) >= 1 and all(lab[n][i] for n in names)
            and len({lab[n][i] for n in names}) == 1 and lab[names[0]][i] != gold[i]]
    log(f"{len(unan)} rows (all {len(names)} annotators labelled the row, same label, != gold)")
    unan.sort(key=lambda i: (key[i]["dataset"], -err.get(i, 0), i))
    for i in unan[: args.n_list]:
        ds = key[i]["dataset"]
        s, t = text.get((ds, key[i]["id"]), ("?", "?"))
        notes = "; ".join(f"{n}: {ann[n][i]['note']}" for n in names if ann[n][i].get("note"))
        log(f"- [{i}] {ds} {key[i]['id']} | gold {gold[i]} | humans {lab[names[0]][i]} | "
            f"{args.llm_name} {llm.get(i, '-')} | model error rate "
            f"{(f'{err[i]:.2f} over {nruns[i]} runs' if i in err else 'n/a')}")
        log(f"    target: {t}")
        log(f"    sentence: {s}")
        if notes:
            log(f"    notes: {notes}")
    log()

    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    with open(args.out, "w", encoding="utf-8") as f:
        f.write("\n".join(OUT) + "\n")
    print(f"written {args.out}")


if __name__ == "__main__":
    main()
