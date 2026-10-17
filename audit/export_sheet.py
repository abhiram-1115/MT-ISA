"""
Export the implicit test rows (restaurant 267 + laptop 175 = 442) to a blind annotation sheet.

Writes
  audit/annotation_sheet.xlsx : sheet "Guideline" + sheet "Annotation" with
                                item | dataset | sentence ([[target]] marked) | target | label | sure | note
                                (label / sure are dropdowns; NO gold, NO instance id)
  audit/key.csv               : item -> dataset, instance_id, gold   (never share this with annotators)

Rows are shuffled with a fixed seed; restaurant rows first, then laptop.
Usage: python audit/export_sheet.py [--seed 13]
"""
import argparse
import csv
import json
import os
import random
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from guideline import DATASETS, GUIDELINE, LABELS, SHEET  # noqa: E402

CODE_CHARS = "ABCDEFGHJKLMNPQRSTUVWXYZ23456789"  # no 0/O, 1/I


def prefix(i):
    return i.rsplit("_aspect_", 1)[0] if "_aspect_" in i else i


def mark_target(row, spans, stats):
    """Wrap the target as [[target]], using the exact span from the *_all file when available."""
    s, t = row["sentence"], row["target"]
    for a, b in spans.get((prefix(row["id"]), t), []):
        if s[a:b] == t:
            stats["span"] += 1
            return s[:a] + "[[" + t + "]]" + s[b:]
    m = re.search(re.escape(t), s, flags=re.IGNORECASE)
    if m:
        stats["first_match"] += 1
        return s[:m.start()] + "[[" + s[m.start():m.end()] + "]]" + s[m.end():]
    stats["not_found"] += 1
    return s + "   (target: [[" + t + "]])"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--seed", type=int, default=13)
    ap.add_argument("--out", default="audit/annotation_sheet.xlsx")
    ap.add_argument("--key", default="audit/key.csv")
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    rng = random.Random(args.seed)
    rows = []
    for ds in ("restaurant", "laptop"):
        impl_path, all_path = DATASETS[ds]
        impl = json.load(open(impl_path, encoding="utf-8"))
        spans = {}
        if os.path.exists(all_path):
            for r in json.load(open(all_path, encoding="utf-8")):
                try:
                    spans.setdefault((str(r["sentence_id"]), r["target"]), []).append((int(r["from"]), int(r["to"])))
                except (KeyError, ValueError, TypeError):
                    pass
        stats = {"span": 0, "first_match": 0, "not_found": 0}
        part = []
        for r in impl:
            if r["gold_polarity"].lower() not in LABELS:
                continue
            part.append({"dataset": ds, "id": r["id"], "gold": r["gold_polarity"].lower(),
                         "sentence": mark_target(r, spans, stats), "target": r["target"]})
        rng.shuffle(part)
        print(f"{ds}: {len(part)} rows | target marked via span {stats['span']}, first match "
              f"{stats['first_match']}, not found {stats['not_found']}")
        rows += part

    codes = set()
    for r in rows:
        while True:
            c = "".join(rng.choice(CODE_CHARS) for _ in range(4))
            if c not in codes:
                codes.add(c)
                r["item"] = c
                break

    wb = Workbook()
    g = wb.active
    g.title = "Guideline"
    g["A1"] = GUIDELINE
    g["A3"] = ("How to fill: in sheet 'Annotation', choose label and sure from the dropdowns for every "
               "row; note is optional. Do not edit, sort away or delete the item column.")
    for cell in ("A1", "A3"):
        g[cell].alignment = Alignment(wrap_text=True, vertical="top")
    g["A1"].font = Font(size=12)
    g.column_dimensions["A"].width = 110
    g.row_dimensions[1].height = 110
    g.row_dimensions[3].height = 35

    ws = wb.create_sheet(SHEET)
    header = ["item", "dataset", "sentence", "target", "label", "sure", "note"]
    ws.append(header)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r["item"], r["dataset"], r["sentence"], r["target"], None, None, None])
    n = len(rows)
    dv_label = DataValidation(type="list", formula1='"positive,negative,neutral"', allow_blank=True)
    dv_sure = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(dv_label)
    ws.add_data_validation(dv_sure)
    dv_label.add(f"E2:E{n + 1}")
    dv_sure.add(f"F2:F{n + 1}")
    for col, w in zip("ABCDEFG", (8, 11, 90, 22, 12, 8, 30)):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(args.out) or ".", exist_ok=True)
    wb.save(args.out)

    with open(args.key, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["item", "dataset", "instance_id", "gold"])
        for r in rows:
            w.writerow([r["item"], r["dataset"], r["id"], r["gold"]])
    print(f"written {args.out} ({n} rows) and {args.key} - keep the key away from annotators")


if __name__ == "__main__":
    main()
